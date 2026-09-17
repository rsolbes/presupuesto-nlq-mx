"""
Cruce entre los dos formatos que publica la SHCP para cada ejercicio.

Un xlsx no puede presentar renglones malformados: el defecto de escapado es
exclusivo del texto delimitado. Lo que este cruce comprueba es si el CSV
publicado PERDIO informacion respecto del xlsx del mismo ejercicio, comparando
renglones utiles y totales por medida.

Uso:    python experiments/cruce_xlsx_csv.py
Salida: experiments/salidas/cruce_xlsx_csv.md
"""

import csv
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl

from perfilado_fuentes import CANONICO, LLAVES, MEDIDAS, a_decimal

RAIZ = Path(__file__).resolve().parents[1]
CRUDO = RAIZ / "data" / "raw"
SALIDA = RAIZ / "experiments" / "salidas" / "cruce_xlsx_csv.md"

EJERCICIOS = ["2020", "2021", "2022", "2023", "2024", "2025"]


def rutas(anio):
    """Devuelve (csv, xlsx) del ejercicio. Tolera el xlsx sin extension."""
    base = "cuenta_publica_" + anio + "_gf_ecd_epe"
    c = CRUDO / (base + ".csv")
    x = CRUDO / (base + ".xlsx")
    if not x.exists() and (CRUDO / base).exists():
        x = CRUDO / base
    return (c if c.exists() else None, x if x.exists() else None)


def acumular(idx, filas, es_xlsx):
    """Recorre filas ya normalizadas a lista y acumula metricas."""
    utiles = 0
    sumas = {m: Decimal(0) for m in MEDIDAS}
    vacios = {m: 0 for m in MEDIDAS}
    mapas = {tuple(k): {} for k, _ in LLAVES}

    for fila in filas:
        utiles += 1
        for m in MEDIDAS:
            v = a_decimal(fila[idx[m]])
            if v is None:
                vacios[m] += 1
            else:
                sumas[m] += v
        for claves, desc in LLAVES:
            k = tuple(fila[idx[c]] for c in claves)
            mapas[tuple(claves)].setdefault(k, set()).add(fila[idx[desc]])

    ambiguas = [
        "+".join(claves)
        for claves, _ in LLAVES
        if max((len(v) for v in mapas[tuple(claves)].values()), default=0) > 1
    ]
    return {"utiles": utiles, "sumas": sumas, "vacios": vacios, "ambiguas": ambiguas,
            "malformados": 0 if es_xlsx else None}


def leer_csv(ruta):
    with open(ruta, encoding="latin-1", newline="") as fh:
        lector = csv.reader(fh)
        crudo = [c.strip() for c in next(lector)]
        ncol = len(crudo)
        nombradas = [c for c in crudo if c]
        idx = {canon: i for i, canon in enumerate(CANONICO[: len(nombradas)])}

        malformados = 0
        buenas = []
        for fila in lector:
            if not fila or not fila[0].strip():
                continue
            if len(fila) != ncol:
                malformados += 1
                continue
            buenas.append(fila)

    r = acumular(idx, buenas, es_xlsx=False)
    r["malformados"] = malformados
    r["encabezado"] = nombradas
    return r


def leer_xlsx(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    crudo = [str(c).strip() if c is not None else "" for c in next(it)]
    nombradas = [c for c in crudo if c]
    idx = {canon: i for i, canon in enumerate(CANONICO[: len(nombradas)])}

    def texto(v):
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)

    buenas = []
    for fila in it:
        if fila is None or fila[0] in (None, ""):
            continue
        buenas.append([texto(v) for v in fila[: len(nombradas)]])
    wb.close()

    r = acumular(idx, buenas, es_xlsx=True)
    r["encabezado"] = nombradas
    return r


def informe(res):
    L = []
    L.append("# Cruce entre CSV y XLSX publicados - Cuenta Publica (SHCP)")
    L.append("")
    L.append("Generado el " + date.today().isoformat()
             + " por `experiments/cruce_xlsx_csv.py`.")
    L.append("")
    L.append("Un xlsx no puede presentar renglones malformados: el defecto de escapado es "
             "exclusivo del texto delimitado. La comparacion util es si el CSV publicado "
             "pierde informacion frente al xlsx del mismo ejercicio.")

    L.append("")
    L.append("## 1. Renglones utiles")
    L.append("")
    L.append("| Ejercicio | CSV | XLSX | Diferencia | Malformados en CSV |")
    L.append("|---|---:|---:|---:|---:|")
    for anio, c, x in res:
        if not c or not x:
            L.append("| {} | {} | {} | - | - |".format(
                anio, "{:,}".format(c["utiles"]) if c else "falta",
                "{:,}".format(x["utiles"]) if x else "falta"))
            continue
        d = c["utiles"] - x["utiles"]
        L.append("| {} | {:,} | {:,} | {:+,} | {:,} |".format(
            anio, c["utiles"], x["utiles"], d, c["malformados"]))

    L.append("")
    L.append("## 2. Totales por medida (millones de pesos)")
    L.append("")
    for anio, c, x in res:
        if not c or not x:
            continue
        L.append("")
        L.append("### Ejercicio " + anio)
        L.append("")
        L.append("| Medida | CSV | XLSX | Diferencia | Coincide |")
        L.append("|---|---:|---:|---:|---|")
        for m in MEDIDAS:
            a = c["sumas"][m] / 1000000
            b = x["sumas"][m] / 1000000
            ok = "si" if abs(a - b) < Decimal("0.05") else "NO"
            L.append("| {} | {:,.1f} | {:,.1f} | {:+,.1f} | {} |".format(
                m, a, b, a - b, ok))

    L.append("")
    L.append("## 3. Celdas de importe sin valor")
    L.append("")
    L.append("| Ejercicio | Formato | " + " | ".join(MEDIDAS) + " |")
    L.append("|---" * (len(MEDIDAS) + 2) + "|")
    for anio, c, x in res:
        for etq, r in (("csv", c), ("xlsx", x)):
            if not r:
                continue
            L.append("| {} | {} | ".format(anio, etq)
                     + " | ".join("{:,}".format(r["vacios"][m]) for m in MEDIDAS) + " |")

    L.append("")
    L.append("## 4. Llaves naturales que resultan ambiguas")
    L.append("")
    L.append("| Ejercicio | Formato | Llaves ambiguas |")
    L.append("|---|---|---|")
    for anio, c, x in res:
        for etq, r in (("csv", c), ("xlsx", x)):
            if not r:
                continue
            L.append("| {} | {} | {} |".format(
                anio, etq, ", ".join("`" + a + "`" for a in r["ambiguas"]) or "ninguna"))

    L.append("")
    L.append("## 5. Encabezados por formato")
    L.append("")
    L.append("| Ejercicio | Formato | Primeras columnas |")
    L.append("|---|---|---|")
    for anio, c, x in res:
        for etq, r in (("csv", c), ("xlsx", x)):
            if not r:
                continue
            L.append("| {} | {} | {} |".format(
                anio, etq, ", ".join("`" + h + "`" for h in r["encabezado"][:6])))

    return "\n".join(L) + "\n"


def main():
    res = []
    for anio in EJERCICIOS:
        rc, rx = rutas(anio)
        print("ejercicio " + anio + " ...", flush=True)
        c = leer_csv(rc) if rc else None
        if rc:
            print("   csv  listo", flush=True)
        x = leer_xlsx(rx) if rx else None
        if rx:
            print("   xlsx listo", flush=True)
        res.append((anio, c, x))
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(informe(res), encoding="utf-8")
    print("")
    print("informe escrito en " + str(SALIDA.relative_to(RAIZ)))


if __name__ == "__main__":
    main()
